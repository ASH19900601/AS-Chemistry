# BuildSubmissionDrafts.py
# Builds a draft Part 2 submission for each group member:
#   Submission_Drafts/<Member>_Part2_Submission.docx  - the school template with the member's
#                                                       code and test plan filled in
#   Submission_Drafts/<Member>_Part2_Submission.md    - the same test plan, a rubric checklist
#                                                       and a plan for the demo video
# Every "Actual outcome" is produced by running the program now, so after changing the code
# run this again:   python Tools/BuildSubmissionDrafts.py
# It uses temporary folders, so the real learning records, reports and backups are not touched.
# This is a group tool, not part of any member's module.

import contextlib
import http.server
import io
import json
import os
import shutil
import sys
import tempfile
import threading
from unittest import mock

ToolFolder = os.path.dirname(os.path.abspath(__file__))
ProgramFolder = os.path.dirname(ToolFolder)
sys.path.insert(0, ProgramFolder)
os.chdir(ProgramFolder)

import AIConnection                # noqa: E402
import AISQGrader                  # noqa: E402
import CorrectionRegrading         # noqa: E402
import DocumentProcessing          # noqa: E402
import GradingFeedback             # noqa: E402
import LearningRecords             # noqa: E402
import PaperDocumentManagement     # noqa: E402
import ReportManagement            # noqa: E402
import SQRecords                   # noqa: E402

OutputFolder = os.path.join(ProgramFolder, 'Submission_Drafts')
TemplateFile = os.path.join(OutputFolder, 'Template', '2_Project_Part_2_Submission_Document.docx')
SystemName = 'AI Structure Question Grader (AI SQ Grader)'
DueDate = 'Week 12, Monday 4:30pm'
Teacher = 'Mr Andrew Trigg'
GoodAnswer2b = ('The rate of reaction increases at a higher temperature because more particles have energy '
                'greater than or equal to the activation energy.')

TempRoot = tempfile.mkdtemp(prefix='SQDrafts')
os.environ['no_proxy'] = '127.0.0.1,localhost'
os.environ['NO_PROXY'] = '127.0.0.1,localhost'


# ===================================================================== helpers

def Run(Function, Inputs, *Arguments):
    # Runs Function with typed inputs; returns (result, printed text).
    Output = io.StringIO()
    with mock.patch('builtins.input', side_effect=list(Inputs)):
        with contextlib.redirect_stdout(Output):
            Result = Function(*Arguments)
    return Result, Output.getvalue()


FolderNames = []


def Clean(Text):
    # Shortens temporary and program paths so the outcome reads like the user's screen.
    for Name in FolderNames:
        Text = Text.replace(os.path.join(TempRoot, Name) + os.sep, '')
    Text = Text.replace(TempRoot + os.sep, '').replace(ProgramFolder + os.sep, '')
    return Text.replace(TempRoot, '').replace(ProgramFolder, '')


def Pick(Printed, *Keys):
    # Returns the printed lines that contain any of the keys, each in quotes, in order.
    Found = []
    for Line in Printed.splitlines():
        Line = Line.strip()
        for Key in Keys:
            if Key in Line and Line not in Found:
                Found.append(Line)
    return ' / '.join('"' + Clean(Line) + '"' for Line in Found)


def FreshFolder(Name):
    Folder = os.path.join(TempRoot, Name)
    FolderNames.append(Name)
    shutil.rmtree(Folder, ignore_errors=True)
    os.makedirs(Folder)
    return Folder


@contextlib.contextmanager
def IsolatedFiles(Name):
    # Points the database, reports, backups and downloads at a new temporary folder.
    Folder = FreshFolder(Name)
    Patches = [mock.patch.object(LearningRecords, 'DatabasePath', os.path.join(Folder, 'LearningRecords.json')),
               mock.patch.object(ReportManagement, 'DatabasePath', os.path.join(Folder, 'LearningRecords.json')),
               mock.patch.object(ReportManagement, 'ReportFolder', os.path.join(Folder, 'Reports')),
               mock.patch.object(ReportManagement, 'BackupFolder', os.path.join(Folder, 'Backups')),
               mock.patch.object(PaperDocumentManagement, 'DownloadFolder', os.path.join(Folder, 'Downloads'))]
    for Patch in Patches:
        Patch.start()
    try:
        yield Folder
    finally:
        for Patch in Patches:
            Patch.stop()


def SampleDocuments():
    Documents = SQRecords.DocumentSet()
    Documents.Paper.Level = 'AS Level'
    Documents.Paper.Subject = 'Chemistry (9701)'
    Documents.Paper.PaperCode = '9701/22/M/J/23'
    Documents.Paper.Session = 'May/June 2023'
    Documents.MarkSchemeFile = os.path.join(ProgramFolder, 'SampleData', 'Sample_Chemistry_MS.txt')
    Documents.AnswerFiles = [os.path.join(ProgramFolder, 'SampleData', 'Sample_Chemistry_Answers.txt'),
                             os.path.join(ProgramFolder, 'SampleData', 'Sample_Answer_Q3.png')]
    Documents.IsComplete = True
    return Documents


def SampleGrading():
    Processed, _ = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), SQRecords.AISettings())
    Grading, Printed = Run(GradingFeedback.GradePaper, [], Processed, SQRecords.AISettings())
    return Grading, Printed


def SampleQuestion(QuestionID, AnswerText):
    Processed, _ = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), SQRecords.AISettings())
    for Question in Processed.Questions:
        if Question.QuestionID == QuestionID:
            Question.AnswerText = AnswerText
            return Question
    raise ValueError(QuestionID)


def Categories(Feedback):
    return ', '.join(Point.PointID + ' ' + Point.Category for Point in Feedback.Points)


class FileServer:
    # A small local web server: serves files from a folder, or answers like an AI service.
    def __init__(self, Folder=None, AIReplies=None):
        Outer = self

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                Name = os.path.join(Folder, os.path.basename(self.path))
                if os.path.isfile(Name):
                    with open(Name, 'rb') as File:
                        Data = File.read()
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(Data)
                else:
                    self.send_error(404)

            def do_POST(self):
                self.rfile.read(int(self.headers['Content-Length']))
                Status, Text = Outer.AIReplies.pop(0)
                self.send_response(Status)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(Text.encode('utf-8'))

            def log_message(self, *Arguments):
                pass

        self.AIReplies = AIReplies or []
        self.Server = http.server.HTTPServer(('127.0.0.1', 0), Handler)
        self.URL = 'http://127.0.0.1:' + str(self.Server.server_port) + '/'
        threading.Thread(target=self.Server.serve_forever, daemon=True).start()

    def Stop(self):
        self.Server.shutdown()
        self.Server.server_close()


def ClosedPortURL():
    Server = http.server.HTTPServer(('127.0.0.1', 0), http.server.BaseHTTPRequestHandler)
    Port = Server.server_port
    Server.server_close()
    return 'http://127.0.0.1:' + str(Port) + '/'


def Completion(Content):
    return json.dumps({'choices': [{'message': {'role': 'assistant', 'content': Content}}]})


def OnlineSettings(Server):
    Settings = SQRecords.AISettings()
    Settings.Provider = Server.URL + 'v1'
    Settings.ModelName = 'test-model'
    Settings.APIKey = 'sk-test'
    return Settings


def Row(Kind, Data, Expected, Actual):
    return {'Data': Kind + ': ' + Data, 'Expected': Expected, 'Actual': Actual}


# ===================================================================== Fox

def FoxTests():
    Tests = []
    Rows = []
    Result, _ = Run(PaperDocumentManagement.SelectPaper, ['2', '2', '22', '2', '2023'])
    Rows.append(Row('Normal', 'AS Level, Chemistry, paper 22, May/June, 2023',
                    'Paper code 9701/22/M/J/23 and session May/June 2023 are stored',
                    'Paper code ' + Result.PaperCode + ', session ' + Result.Session))
    Result, Printed = Run(PaperDocumentManagement.SelectPaper, ['2', '2', '27', '2a', '22', '2', '2023'])
    Rows.append(Row('Abnormal', 'paper number 27, then 2a, then 22',
                    'Error message for 27 and 2a; the number is asked for again until 22 is typed',
                    Pick(Printed, 'A paper number is two digits') + ' shown ' + str(Printed.count('A paper number is two digits')) + ' times; then ' + Result.PaperCode + ' stored'))
    Result, Printed = Run(PaperDocumentManagement.SelectPaper, ['2', '2', '22', '2', '2015', '2016'])
    Rows.append(Row('Boundary', 'year 2015, then 2016 (the first year allowed)',
                    '2015 rejected with the allowed range; 2016 accepted',
                    Pick(Printed, 'The year must be from') + '; then session "' + Result.Session + '" stored'))
    Result, Printed = Run(PaperDocumentManagement.SelectPaper, ['5', 'two', '2', '2', '22', '2', '2023'])
    Rows.append(Row('Abnormal', 'level menu: 5, then "two"', 'Error message each time; menu shown again',
                    Pick(Printed, 'That number is not in the list', 'Please type a whole number')))
    Tests.append({'Name': 'Select the paper', 'Purpose': 'To check that the paper details are only accepted in the correct form and are turned into the Cambridge paper code.', 'Rows': Rows})

    Rows = []
    Inputs = ['SampleData/Sample_Chemistry_Answers.txt', 'Answers.docx', 'Missing.txt', 'SampleData/Sample_Chemistry_Answers.txt', '']
    Result, Printed = Run(PaperDocumentManagement.InputAnswerFiles, Inputs)
    Rows.append(Row('Normal', 'SampleData/Sample_Chemistry_Answers.txt', 'File added to the list', Pick(Printed, 'Added:')))
    Rows.append(Row('Abnormal', 'Answers.docx (Word file)', 'Rejected: file type not supported', Pick(Printed, 'not supported')))
    Rows.append(Row('Abnormal', 'Missing.txt (does not exist)', 'Rejected: file cannot be found', Pick(Printed, 'cannot be found')))
    Rows.append(Row('Abnormal', 'the same answer file a second time', 'Rejected: already in the list',
                    Pick(Printed, 'already in the list') + '; list holds ' + str(len(Result)) + ' file'))
    Tests.append({'Name': 'Upload answer files', 'Purpose': 'To check that only supported files that exist are added, and each file only once.', 'Rows': Rows})

    Rows = []
    Inputs = ['2', '2', '22', '2', '2023', '', '']
    Result, Printed = Run(PaperDocumentManagement.ManageDocuments, Inputs, SQRecords.DownloadSettings())
    Rows.append(Row('Abnormal', 'paper selected; Enter pressed for the answer file and for the mark scheme',
                    'Each missing item is listed and marking does not begin (IsComplete = FALSE)',
                    Pick(Printed, 'Missing:', 'Marking cannot begin') + ' (IsComplete = ' + str(Result.IsComplete).upper() + ')'))
    with mock.patch.object(AISQGrader.DocumentProcessing, 'ProcessDocuments', wraps=DocumentProcessing.ProcessDocuments) as Processing:
        _, Printed = Run(AISQGrader.RunMarkingSession, Inputs, SQRecords.AISettings(), SQRecords.DownloadSettings())
    Rows.append(Row('Abnormal', 'the same inputs through the main module (AI SQ Grader)',
                    'The session stops before the Document Processing Module is called',
                    Pick(Printed, 'session has stopped') + '; Document Processing called ' + str(Processing.call_count) + ' times'))
    Inputs = ['2', '2', '22', '2', '2023', 'SampleData/Sample_Chemistry_Answers.txt', 'SampleData/Sample_Answer_Q3.png', '', 'SampleData/Sample_Chemistry_MS.txt', '']
    Result, Printed = Run(PaperDocumentManagement.ManageDocuments, Inputs, SQRecords.DownloadSettings())
    Processed, Printed2 = Run(DocumentProcessing.ProcessDocuments, [], Result, SQRecords.AISettings())
    Rows.append(Row('Normal', 'paper, Sample_Chemistry_Answers.txt, Sample_Answer_Q3.png and Sample_Chemistry_MS.txt',
                    'IsComplete = TRUE; the DocumentSet is read by the Document Processing Module (integration)',
                    Pick(Printed, 'All documents are ready') + ' (IsComplete = ' + str(Result.IsComplete).upper() + '); Document Processing: ' + Pick(Printed2, 'Mark scheme read')))
    Tests.append({'Name': 'Missing uploads', 'Purpose': 'To check that marking only begins when the paper details, the answers and the mark scheme are all available.', 'Rows': Rows})

    Rows = []
    with IsolatedFiles('FoxDownload'):
        Folder = FreshFolder('FoxServer')
        with open(os.path.join(Folder, '9701_s23_ms_22.pdf'), 'wb') as File:
            File.write(b'%PDF-1.4\n% sample mark scheme for testing\n%%EOF\n')
        with open(os.path.join(Folder, '9701_s23_ms_21.pdf'), 'wb') as File:
            File.write(b'<html><body>Please log in</body></html>')
        Server = FileServer(Folder)
        Download = SQRecords.DownloadSettings()
        Download.Enabled = True
        Download.ServiceURL = Server.URL
        Paper = SQRecords.PaperDetails()
        for Code, Kind, Data, Expected in [
                ('9701/22/M/J/23', 'Normal', 'paper 9701/22/M/J/23; the test server has the PDF', 'PDF downloaded and saved in the Downloads folder'),
                ('9701/21/M/J/23', 'Abnormal', 'paper 9701/21/M/J/23; the server sends a web page instead of a PDF', 'Rejected because the file is not a PDF'),
                ('9701/23/M/J/23', 'Abnormal', 'paper 9701/23/M/J/23; the file is not on the server', 'Download failed message with the HTTP error number')]:
            Paper.PaperCode = Code
            _, Printed = Run(PaperDocumentManagement.DownloadPaperFile, [], Paper, 'ms', Download)
            Rows.append(Row(Kind, Data, Expected, Pick(Printed, 'saved as', 'Download failed')))
        Server.Stop()
        Download.ServiceURL = ClosedPortURL()
        Paper.PaperCode = '9701/22/M/J/23'
        _, Printed = Run(PaperDocumentManagement.DownloadPaperFile, [], Paper, 'ms', Download)
        Rows.append(Row('Abnormal', 'the download server is switched off', 'Message that the service could not be reached', Pick(Printed, 'Download failed')))
    Tests.append({'Name': 'Download the mark scheme', 'Purpose': 'To check that a mark scheme can be downloaded from the online service and that anything that is not a PDF is rejected.', 'Rows': Rows})

    Rows = []
    Settings, Printed = Run(AISQGrader.LoadAISettings, [], [['Provider', 'Offline']])
    _, Printed = Run(AISQGrader.DisplaySettings, [], Settings, SQRecords.DownloadSettings())
    Rows.append(Row('Normal', 'Settings.txt contains Provider=Offline', 'Offline demo marker selected; no API key is asked for', Pick(Printed, 'AI service')))
    with mock.patch.dict(os.environ, {'SQ_API_KEY': ''}):
        with mock.patch('getpass.getpass', return_value=''):
            Settings, Printed = Run(AISQGrader.LoadAISettings, [], [['Provider', 'OpenAI'], ['ModelName', 'gpt-4o-mini']])
    Rows.append(Row('Abnormal', 'Provider=OpenAI, but Enter is pressed instead of an API key',
                    'Message; the offline demo marker is used instead', Pick(Printed, 'No API key') + ' (Provider = ' + Settings.Provider + ')'))
    _, Printed = Run(AISQGrader.ReadSettingsFile, [], 'NoSettings.txt')
    Rows.append(Row('Abnormal', 'the settings file is missing', 'Message; default settings are used', Pick(Printed, 'Settings file not found')))
    Tests.append({'Name': 'Load the settings (main module)', 'Purpose': 'To check that the main module reads its settings file, never needs the API key to be saved, and still runs when a setting is missing.', 'Rows': Rows})
    Rows = []
    Details = SQRecords.PaperDetails()
    Result, Printed = Run(PaperDocumentManagement.InputQuestionPaper, [''], Details, SQRecords.DownloadSettings())
    Rows.append(Row('Normal', 'Enter pressed (no question paper)', 'Skipped without an error; marking can still begin',
                    'QuestionPaperFile = "' + Result + '"'))
    Result, Printed = Run(PaperDocumentManagement.InputQuestionPaper, ['Paper.docx', 'SampleData/Sample_Answer_Q3.png'], Details, SQRecords.DownloadSettings())
    Rows.append(Row('Abnormal', 'Paper.docx, then a photo of the paper (.png)', 'Word file rejected; photo accepted',
                    Pick(Printed, 'Not accepted') + '; then "' + os.path.basename(Result) + '" accepted'))
    Tests.append({'Name': 'Upload the question paper (optional)', 'Purpose': 'To check that the question paper can be added to help the AI, but is not required.', 'Rows': Rows})

    Rows = []
    Server = FileServer(AIReplies=[(200, Completion('{"ok": true}')), (200, Completion('I am a chatbot.')), (401, '{"error": "invalid key"}')])
    Settings = OnlineSettings(Server)
    for Kind, Data, Expected in [('Normal', 'a test server that answers like an OpenAI-compatible AI service', 'Message that the AI service is working'),
                                 ('Abnormal', 'the server answers with plain text instead of JSON', 'Message to choose another model'),
                                 ('Abnormal', 'the server answers HTTP 401 (wrong API key)', 'Message that the test failed, with the error')]:
        Result, Printed = Run(AIConnection.TestConnection, [], Settings)
        Rows.append(Row(Kind, Data, Expected, Pick(Printed, 'working', 'not with JSON', 'failed') + ' (returned ' + str(Result).upper() + ')'))
    Server.Stop()
    Result, Printed = Run(AIConnection.TestConnection, [], SQRecords.AISettings())
    Rows.append(Row('Boundary', 'the offline demo marker is chosen', 'Message that there is no AI service to test', Pick(Printed, 'offline demo marker')))
    Tests.append({'Name': 'Test the AI connection', 'Purpose': 'To check that the program reports clearly whether the chosen AI service, model and API key work before a paper is marked.', 'Rows': Rows})
    return Tests


# ===================================================================== Kevin

def KevinTests():
    Tests = []
    Rows = []
    Questions, Printed = Run(DocumentProcessing.ParseMarkScheme, [], DocumentProcessing.ReadTextFile(SampleDocuments().MarkSchemeFile))
    Question = Questions[2]
    Rows.append(Row('Normal', 'SampleData/Sample_Chemistry_MS.txt',
                    '5 questions; question 2(a) has 3 marks, 3 marking points and the rule P3 DEPENDS ON P2',
                    str(len(Questions)) + ' questions (' + ', '.join(Q.QuestionID for Q in Questions) + '); 2(a): '
                    + str(Question.MarkingRules.MaxMark) + ' marks, ' + str(len(Question.MarkingPoints)) + ' points, rule "' + Question.MarkingRules.RuleNotes + '"'))
    _, Printed = Run(DocumentProcessing.ParseMarkScheme, [], ['QUESTION 1 [1]', 'QUESTION 4 two', 'P1 [1] x'])
    Rows.append(Row('Abnormal', 'line 2 is "QUESTION 4 two" (marks missing)', 'The line is ignored with a message giving its line number', Pick(Printed, 'line 2')))
    _, Printed = Run(DocumentProcessing.ParseMarkScheme, [], ['QUESTION 6 [1]', 'QUESTION 7 [1]', 'P1 [1] x'])
    Rows.append(Row('Abnormal', 'QUESTION 6 [1] with no marking points', 'Question 6 is left out with a message', Pick(Printed, 'Question 6')))
    _, Printed = Run(DocumentProcessing.ParseMarkScheme, [], ['QUESTION 5 [2]', 'P1 [1] only one point'])
    Rows.append(Row('Boundary', 'QUESTION 5 [2] whose marking points add up to only 1', 'Warning, but the question is kept', Pick(Printed, 'Warning')))
    Tests.append({'Name': 'Read the mark scheme', 'Purpose': 'To check that question numbers, marking points, marks and marking rules are extracted, and badly written lines are reported.', 'Rows': Rows})

    Rows = []
    Processed, Printed = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), SQRecords.AISettings())
    Grading, Printed2 = Run(GradingFeedback.GradePaper, [], Processed, SQRecords.AISettings())
    Q3 = Grading.QuestionResults[4]
    Rows.append(Row('Abnormal', 'question 3 is answered only in the photo Sample_Answer_Q3.png',
                    'Question 3 listed as uncertain and flagged for review; its points are "Not assessed" (not marked wrong)',
                    Pick(Printed, 'Question 3:') + '; ' + Pick(Printed2, 'flagged for review', 'needs review') + '; ' + Categories(Q3.Feedback)))
    Rows.append(Row('Normal', 'answers 1(a), 1(b), 2(a) and 2(b) labelled in Sample_Chemistry_Answers.txt',
                    'The four answers are matched to their questions', str(len([Line for Line in Printed.splitlines() if Line.strip().endswith('found')])) + ' questions shown as "found"; uncertain matches: ' + str(Processed.UncertainMatches)))
    Folder = FreshFolder('KevinAnswers')
    FileName = os.path.join(Folder, 'Answers.txt')
    with open(FileName, 'w', encoding='utf-8') as File:
        File.write('1(a) protons and neutrons are in the nucleus\n9(z) an answer to a question that is not in the paper\n')
    Documents = SampleDocuments()
    Documents.AnswerFiles = [FileName]
    _, Printed = Run(DocumentProcessing.ProcessDocuments, [], Documents, SQRecords.AISettings())
    Rows.append(Row('Abnormal', 'an answer labelled 9(z), which is not in the mark scheme', 'Message; 9(z) is listed as an uncertain match', Pick(Printed, '9(z)')))
    Tests.append({'Name': 'Unclear question-answer match', 'Purpose': 'To check that answers that cannot be matched with certainty are flagged for review instead of being marked.', 'Rows': Rows})

    Rows = []
    Settings = SQRecords.AISettings()
    Question = SampleQuestion('2(b)', GoodAnswer2b)
    Feedback, _ = Run(GradingFeedback.AssessMarkingPoints, [], Question, True, Settings)
    Rows.append(Row('Normal', '2(b): "' + GoodAnswer2b + '"', 'P1 and P2 Clearly correct; estimated mark 2 / 2',
                    Categories(Feedback) + '; estimated mark ' + str(GradingFeedback.CalculateEstimatedMark(Feedback, Question.MarkingRules)) + ' / 2'))
    Result = Grading.QuestionResults[2]
    _, Printed = Run(GradingFeedback.DisplayQuestionResult, [], Result)
    Rows.append(Row('Normal', '2(a) sample answer (the 1 : 1 ratio is not explained)',
                    'P2 Possibly correct; P3 is not counted because it depends on P2; estimated mark 1 / 3',
                    Categories(Result.Feedback) + '; ' + Pick(Printed, 'Not counted', 'Estimated mark')))
    Result = Grading.QuestionResults[3]
    _, Printed = Run(GradingFeedback.DisplayQuestionResult, [], Result)
    Rows.append(Row('Abnormal', '2(b) sample answer: "The rate gets faster because the particles collide more often. Catalysts are not used up in the reaction."',
                    'Both points Incorrect, 0 / 2; a reason is given for each lost mark; the catalyst sentence is listed as outside the mark scheme',
                    Categories(Result.Feedback) + '; ' + Pick(Printed, 'Estimated mark', 'Likely reason', 'Outside the mark scheme')))
    Rules = SQRecords.MarkRules()
    Rules.MaxMark = 2
    Feedback = SQRecords.PointFeedbackSet()
    for ID in ['P1', 'P2', 'P3']:
        Point = SQRecords.PointFeedback()
        Point.PointID = ID
        Point.Marks = 1
        Point.Category = SQRecords.CategoryClearlyCorrect
        Feedback.Points.append(Point)
    Rows.append(Row('Boundary', 'three Clearly correct points worth 1 mark each; maximum mark 2', 'Estimated mark capped at 2',
                    'Estimated mark ' + str(GradingFeedback.CalculateEstimatedMark(Feedback, Rules))))
    Tests.append({'Name': 'Assess marking points and estimate the mark', 'Purpose': 'To check the four feedback categories and that only Clearly correct points count, following the marking rules.', 'Rows': Rows})

    Rows = []
    Points = SampleQuestion('2(a)', '').MarkingPoints
    for Kind, Data, Expected, Text in [
            ('Normal', 'AI reply {"points": [P1 Clearly correct, P2 Incorrect, P3 Possibly correct]}', 'Categories read into the PointFeedbackSet; NeedsReview = FALSE',
             '{"points": [{"point_id": "P1", "category": "Clearly correct"}, {"point_id": "P2", "category": "Incorrect"}, {"point_id": "P3", "category": "Possibly correct"}]}'),
            ('Abnormal', 'AI reply "Sorry, I cannot help with that."', 'Every point "Not assessed"; NeedsReview = TRUE; no marks given', 'Sorry, I cannot help with that.'),
            ('Abnormal', 'AI reply with the category "Partly right" for P1', 'P1 "Not assessed" with a message; NeedsReview = TRUE',
             '{"points": [{"point_id": "P1", "category": "Partly right"}, {"point_id": "P2", "category": "Incorrect"}, {"point_id": "P3", "category": "Incorrect"}]}')]:
        Response = SQRecords.AIResponse()
        Response.Success = True
        Response.ReplyText = Text
        Feedback = GradingFeedback.ParsePointFeedback(Response, Points)
        Rows.append(Row(Kind, Data, Expected, Categories(Feedback) + '; NeedsReview = ' + str(Feedback.NeedsReview).upper()
                        + ('; "' + Feedback.Points[0].Explanation + '"' if Feedback.Points[0].Category == SQRecords.CategoryNotAssessed else '')))
    Tests.append({'Name': 'Check the AI reply', 'Purpose': 'To check that a reply that is not valid JSON, or has an unknown category, can never give marks.', 'Rows': Rows})

    Rows = []
    Question = SampleQuestion('1(a)', 'The nucleus contains protons and neutrons.')
    Reply = json.dumps({'points': [{'point_id': 'P1', 'category': 'Clearly correct', 'explanation': 'Both particles are named.', 'evidence': 'protons and neutrons'},
                                   {'point_id': 'P2', 'category': 'Incorrect', 'explanation': 'Electron shells are not mentioned.', 'evidence': ''}], 'outside_ms': []})
    Server = FileServer(AIReplies=[(200, Completion(Reply)), (401, '{"error": "invalid key"}')])
    Settings = SQRecords.AISettings()
    Settings.Provider = Server.URL + 'v1'
    Settings.ModelName = 'test-model'
    Settings.APIKey = 'sk-test'
    Feedback, Printed = Run(GradingFeedback.CompareWithMarkingPoints, [], Question, Settings)
    Rows.append(Row('Normal', 'question 1(a) sent to a test server that answers like an OpenAI-compatible AI service',
                    'The AI judgement is used: P1 Clearly correct, P2 Incorrect', Categories(Feedback) + '; P2 explanation "' + Feedback.Points[1].Explanation + '"'))
    Feedback, Printed = Run(GradingFeedback.CompareWithMarkingPoints, [], Question, Settings)
    Rows.append(Row('Abnormal', 'the test server answers HTTP 401 (wrong API key)', 'Error message; points "Not assessed", NeedsReview = TRUE (the question is not marked wrong)',
                    Pick(Printed, 'failed') + '; ' + Categories(Feedback) + '; NeedsReview = ' + str(Feedback.NeedsReview).upper()))
    Server.Stop()
    Settings.Provider = ClosedPortURL() + 'v1'
    Feedback, Printed = Run(GradingFeedback.CompareWithMarkingPoints, [], Question, Settings)
    Rows.append(Row('Abnormal', 'the AI server is switched off', 'After 3 attempts: network error message; points "Not assessed"', Pick(Printed, 'failed') + '; ' + Categories(Feedback)))
    Tests.append({'Name': 'Request an AI judgement', 'Purpose': 'To check that the module works with an online AI service and that a failed request flags the question instead of marking it wrong.', 'Rows': Rows})
    Rows = []
    Folder = FreshFolder('KevinMarkScheme')
    FileName = os.path.join(Folder, '9701_s23_ms_22.pdf')
    import pymupdf
    Document = pymupdf.open()
    Document.new_page().insert_text((72, 72), 'Cambridge International AS & A Level  MARK SCHEME  Maximum Mark: 60  Published')
    Page = Document.new_page()
    Lines = ['Question   Answer   Marks', '1(a)  protons and neutrons in nucleus (1); electrons in shells / energy levels (1)   2',
             '1(b)  same number of protons, different number of neutrons   1',
             '2(b)  rate increases (1); more particles have E >= Ea (1) dep on first mark   2']
    for Index in range(0, len(Lines)):
        Page.insert_text((40, 72 + 16 * Index), Lines[Index], fontsize=9)
    Document.save(FileName)
    Document.close()
    Reply = {'questions': [
        {'question': '1(a)', 'max_mark': 2, 'points': [{'id': 'P1', 'marks': 1, 'requirement': 'protons and neutrons in the nucleus'},
                                                       {'id': 'P2', 'marks': 1, 'requirement': 'electrons in shells / energy levels'}], 'rules': []},
        {'question': '1(b)', 'max_mark': 1, 'points': [{'id': 'P1', 'marks': 1, 'requirement': 'same number of protons, different number of neutrons'}], 'rules': []},
        {'question': '2(b)', 'max_mark': 2, 'points': [{'id': 'P1', 'marks': 1, 'requirement': 'rate increases'},
                                                       {'id': 'P2', 'marks': 1, 'requirement': 'more particles have energy >= activation energy'}], 'rules': ['P2 DEPENDS ON P1']}]}
    Server = FileServer(AIReplies=[(200, Completion(json.dumps(Reply))), (401, '{"error": "invalid key"}')])
    Settings = OnlineSettings(Server)
    with mock.patch.object(DocumentProcessing, 'ParsedFolder', os.path.join(Folder, 'ParsedMarkSchemes')):
        Questions, Printed = Run(DocumentProcessing.ReadMarkScheme, [], FileName, Settings)
        Rows.append(Row('Normal', 'a 2-page mark scheme PDF in the Cambridge layout (cover page, then a table of answers); test AI server',
                        'Cover page skipped; marking points read by the AI and saved in the SQ text format for checking',
                        Pick(Printed, 'Skipped', 'saved as') + '; ' + str(len(Questions)) + ' questions (' + ', '.join(Q.QuestionID for Q in Questions)
                        + '); 2(b) rule "' + Questions[2].MarkingRules.RuleNotes + '"'))
        Questions, Printed = Run(DocumentProcessing.ReadMarkScheme, [], FileName, Settings)
        Rows.append(Row('Abnormal', 'the same PDF, but the AI service refuses the request (HTTP 401)', 'Marking stops, so no incomplete mark scheme is used',
                        Pick(Printed, 'could not read part', 'Marking has stopped') + ' (' + str(len(Questions)) + ' questions)'))
    Server.Stop()
    Questions, Printed = Run(DocumentProcessing.ReadMarkScheme, [], FileName, SQRecords.AISettings())
    Rows.append(Row('Abnormal', 'the same PDF with the offline demo marker', 'Message that an AI service or the SQ text format is needed',
                    Pick(Printed, 'not in the SQ text format')))
    Tests.append({'Name': 'Read a Cambridge mark scheme with AI', 'Purpose': 'To check that a mark scheme PDF is turned into marking points, and that marking stops if part of it cannot be read.', 'Rows': Rows})

    Rows = []
    Server = FileServer(AIReplies=[
        (200, Completion(json.dumps({'answers': [{'question': '3', 'answer': 'A covalent bond is a pair of electrons shared between two atoms. The nuclei are attracted to the shared electrons.', 'confidence': 'high'}]}))),
        (200, Completion(json.dumps({'answers': [{'question': '3', 'answer': 'A covalent ... (smudged)', 'confidence': 'low'}]}))),
        (400, '{"error": "image input is not supported by this model"}')])
    Settings = OnlineSettings(Server)
    Processed, Printed = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), Settings)
    Rows.append(Row('Normal', 'Sample_Answer_Q3.png (photo of handwritten answer 3); test AI server that reads images',
                    'Answer 3 read from the photo and matched to question 3, not flagged',
                    Pick(Printed, 'answer(s) read from', 'found (read from handwriting)') + '; text read: "' + Processed.Questions[4].AnswerText[0:48] + '..."'))
    Processed, Printed = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), Settings)
    Rows.append(Row('Abnormal', 'the AI reads the photo with low confidence', 'Question 3 flagged for review instead of being marked',
                    Pick(Printed, 'could not be read with confidence', 'UNCERTAIN')))
    Processed, Printed = Run(DocumentProcessing.ProcessDocuments, [], SampleDocuments(), Settings)
    Rows.append(Row('Abnormal', 'the chosen model does not accept images (HTTP 400)', 'Message; question 3 flagged for review',
                    Pick(Printed, 'Could not read', 'UNCERTAIN')))
    Server.Stop()
    Tests.append({'Name': 'Read handwriting from a photo', 'Purpose': 'To check that handwritten answers in photos and scans are read by the AI, and that an unclear reading is flagged instead of marked.', 'Rows': Rows})
    return Tests


# ===================================================================== Caroline

def CarolineTests():
    Tests = []
    Grading, _ = SampleGrading()
    Settings = SQRecords.AISettings()
    Rows = []
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '1', GoodAnswer2b, ''], Grading, Settings)
    Rows.append(Row('Normal', 'question 2(b), typed answer "' + GoodAnswer2b + '"',
                    'Reassessed with the same marking points by the Grading and Feedback Module: 0 / 2 -> 2 / 2; changes listed',
                    Pick(Printed, 'Estimated mark:', 'P1: Incorrect', 'P2: Incorrect')))
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '2', 'SampleData/Sample_Corrected_2b.txt'], Grading, Settings)
    Rows.append(Row('Normal', 'question 2(b), answer loaded from SampleData/Sample_Corrected_2b.txt', 'Same result as the typed answer: 0 / 2 -> 2 / 2',
                    Pick(Printed, 'Estimated mark:') + '; attempts ' + str([A.AttemptNo for A in Result.Attempts])))
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '1', 'The rate of reaction increases at a higher temperature.', ''], Grading, Settings)
    Rows.append(Row('Normal', 'question 2(b), partly corrected: "The rate of reaction increases at a higher temperature."', '0 / 2 -> 1 / 2; P2 still to improve',
                    Pick(Printed, 'Estimated mark:', 'Still to improve')))
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '3', 'SampleData/Sample_Answer_Q3.png'], Grading, Settings)
    Rows.append(Row('Abnormal', 'question 2(b), photo of a correction (Sample_Answer_Q3.png) with the offline demo marker', 'Flagged for teacher review; mark not changed',
                    Pick(Printed, 'flagged for review', 'Estimated mark:') + '; NeedsReview = ' + str(Result.Attempts[1].Feedback.NeedsReview).upper()))
    Server = FileServer(AIReplies=[
        (200, Completion(json.dumps({'answers': [{'question': '2(b)', 'answer': GoodAnswer2b, 'confidence': 'high'}]}))),
        (200, Completion(json.dumps({'points': [{'point_id': 'P1', 'category': 'Clearly correct', 'explanation': 'States that the rate increases.'},
                                                {'point_id': 'P2', 'category': 'Clearly correct', 'explanation': 'Refers to particles with at least the activation energy.'}],
                                     'outside_ms': []})))])
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '3', 'SampleData/Sample_Answer_Q3.png'], Grading, OnlineSettings(Server))
    Server.Stop()
    Rows.append(Row('Normal', 'question 2(b), photo of a handwritten correction; test AI server that reads images',
                    'The photo is read by the AI (Document Processing Module) and the correction is marked: 0 / 2 -> 2 / 2',
                    Pick(Printed, 'Read from the photo', 'Estimated mark:')))
    Tests.append({'Name': 'Submit a corrected response', 'Purpose': 'To check that a correction becomes a new attempt, marked against the original criteria, and is compared with the earlier attempt.', 'Rows': Rows})

    Rows = []
    _, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '7', '2(b)', '1', GoodAnswer2b, ''], Grading, Settings)
    Rows.append(Row('Abnormal', 'question number 7 (not in the paper)', 'Error message; the question number is asked for again', Pick(Printed, 'Question 7')))
    _, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['maybe', 'N'], Grading, Settings)
    Rows.append(Row('Abnormal', '"maybe" for the Y/N question', 'Error message; asked again', Pick(Printed, 'Please type Y or N')))
    Result, Printed = Run(CorrectionRegrading.CorrectAndRegrade, ['Y', '2(b)', '1', ''], Grading, Settings)
    Rows.append(Row('Boundary', 'Enter pressed straight away, so the corrected answer is empty', 'No correction submitted (HasCorrection = FALSE)',
                    Pick(Printed, 'answer is empty') + ' (HasCorrection = ' + str(Result.HasCorrection).upper() + ')'))
    Tests.append({'Name': 'Choose the question to correct', 'Purpose': 'To check that invalid choices are rejected with clear messages.', 'Rows': Rows})

    Rows = []
    with IsolatedFiles('CarolineRecords') as Folder:
        Regrade = CorrectionRegrading.RegradeQuestion(Grading, 3, GoodAnswer2b, '', Settings)
        Learning, Printed = Run(LearningRecords.UpdateLearningRecords, ['6'], Grading, Regrade)
        Learning, Printed = Run(LearningRecords.UpdateLearningRecords, ['6'], SQRecords.GradingResult(), SQRecords.RegradeResult())
        Rows.append(Row('Normal', 'sample paper marked, 2(b) corrected and saved; then the module is started again with no new results ("close and reopen")',
                        'The saved paper score, correction and mistake notebook are shown again',
                        Pick(Printed, '9701/22/M/J/23', 'Correction attempts saved', 'Mistake notebook entries')))
        Progress = Learning.Progress
        Rows.append(Row('Normal', 'the same saved records', 'Original score shown apart from corrections; reasons for lost marks from most to least frequent',
                        Pick(Printed, 'Original paper results', 'Questions now fully corrected', 'Incomplete explanation (', 'Calculation mistake (')))
        with open(os.path.join(Folder, 'LearningRecords.json'), 'w', encoding='utf-8') as File:
            File.write('{ damaged')
        _, Printed = Run(LearningRecords.UpdateLearningRecords, [], SQRecords.GradingResult(), SQRecords.RegradeResult())
        Rows.append(Row('Abnormal', 'the records file has been damaged (it is not valid JSON)', 'Message; the damaged file is kept and a new record is started',
                        Pick(Printed, 'damaged', 'No papers have been marked yet')))
    Tests.append({'Name': 'Save and reopen the learning records', 'Purpose': 'To check that results, attempts and the mistake notebook remain available after the application is closed and reopened.', 'Rows': Rows})

    Rows = []
    with IsolatedFiles('CarolineNotebook'):
        Regrade = CorrectionRegrading.RegradeQuestion(Grading, 3, GoodAnswer2b, '', Settings)
        _, Printed = Run(LearningRecords.UpdateLearningRecords, ['2', '1', '4', '2', '6'], Grading, Regrade)
        Parts = Printed.split('Mistake notebook - filter by:')
        Rows.append(Row('Normal', 'filter by topic: "Atomic structure"', 'Questions 1(a) and 3 only', Pick(Parts[0].split('Choose the topic:')[1], 'question')))
        Rows.append(Row('Normal', 'filter by correction status: "Corrected"', 'Question 2(b) with both attempts (mark 0, then 2)', Pick(Parts[1], 'question', 'Attempt')))
        Database, _ = Run(LearningRecords.LoadDatabase, [], LearningRecords.DatabasePath)
        Rows.append(Row('Normal', 'filter by subject: "Chemistry (9701)"', 'All 4 entries of the paper',
                        str(len(LearningRecords.FilterMistakes(Database, 0, 'Chemistry (9701)'))) + ' entries'))
        Rows.append(Row('Abnormal', 'filter by paper: "9702/22/M/J/23" (no mistakes saved for this paper)', 'Message that no entries match',
                        Pick(Run(LearningRecords.DisplayMistakes, [], Database, LearningRecords.FilterMistakes(Database, 2, '9702/22/M/J/23'))[1], 'No mistake')))
    Tests.append({'Name': 'Filter the mistake notebook', 'Purpose': 'To check that the mistake notebook can be filtered by subject, topic, paper and correction status.', 'Rows': Rows})
    return Tests


# ===================================================================== Andy

def AndyTests():
    Tests = []
    Grading, _ = SampleGrading()
    Settings = SQRecords.AISettings()
    Rows = []
    with IsolatedFiles('AndyReport') as Folder:
        Regrade = CorrectionRegrading.RegradeQuestion(Grading, 3, GoodAnswer2b, '', Settings)
        Learning, _ = Run(LearningRecords.UpdateLearningRecords, ['6'], Grading, Regrade)
        Outcome, Printed = Run(ReportManagement.ManageReports, ['N', 'N'], Grading, Learning)
        Detail = ''
        try:
            import pymupdf
            Document = pymupdf.open(Outcome.ReportFile)
            Text = ''.join(Page.get_text() for Page in Document)
            Detail = '; the PDF opens with ' + str(Document.page_count) + ' pages and contains ' + ', '.join(
                '"' + Key + '"' for Key in ['Estimated total: 3 / 10 (30%)', 'Likely reason for lost mark', 'Attempt 2: 2 / 2 (correction)'] if Key in Text)
            Document.close()
        except ImportError:
            pass
        Rows.append(Row('Normal', 'GradingResult of the sample paper (from the Grading and Feedback Module) and LearningData (from the Learning Records Module)',
                        'PDF report saved with the paper details, total, marking-point explanations, error diagnoses and attempt history',
                        Pick(Printed, 'PDF report saved') + Detail))
        Special = 'H' + chr(0x2082) + 'O ' + chr(0x21cc) + ' H' + chr(0x207a) + ' + OH' + chr(0x207b) + ', ' + chr(0x0394) + 'H = -57 kJ (' + chr(0x6e29) + chr(0x5ea6) + ')'
        Safe = ReportManagement.MakePdfSafe(Special)
        SpecialFile = os.path.join(Folder, 'Special.pdf')
        Written = ReportManagement.WritePdfReport(['## Special characters', Special], SpecialFile)
        with open(SpecialFile, 'rb') as File:
            SpecialData = File.read()
        Rows.append(Row('Abnormal', 'answer text with characters the PDF font cannot show: "' + Special + '"',
                        'Subscripts and superscripts written as plain digits and signs, the arrow as "<=>", Delta by name, characters with no font shape as "?", brackets escaped; the PDF is still written',
                        'Written to the PDF as "' + Safe + '"; WritePdfReport returned ' + str(Written).upper()
                        + '; file starts with %PDF and ends with %%EOF: ' + str(SpecialData.startswith(b'%PDF') and SpecialData.rstrip().endswith(b'%%EOF')).upper()))
        Outcome, Printed = Run(ReportManagement.ManageReports, ['N', 'N'], SQRecords.GradingResult(), Learning)
        Rows.append(Row('Boundary', 'no marked questions (records opened from the main menu)', 'No report is created',
                        'ReportFile = "' + Outcome.ReportFile + '"; no "PDF report saved" message: ' + str('PDF report saved' not in Printed).upper()))
    Tests.append({'Name': 'Create the PDF report', 'Purpose': 'To check that the report shows the paper details, estimated scores, marking-point explanations and error diagnoses.', 'Rows': Rows})

    Rows = []
    with IsolatedFiles('AndyBackup') as Folder:
        Outcome, Printed = Run(ReportManagement.ManageReports, ['Y', 'N'], SQRecords.GradingResult(), SQRecords.LearningData())
        Rows.append(Row('Abnormal', 'export a backup before any records have been saved', 'Message that there is nothing to back up', Pick(Printed, 'no saved learning records')))
        Regrade = CorrectionRegrading.RegradeQuestion(Grading, 3, GoodAnswer2b, '', Settings)
        Learning, _ = Run(LearningRecords.UpdateLearningRecords, ['6'], Grading, Regrade)
        Outcome, Printed = Run(ReportManagement.ManageReports, ['Y', 'N'], SQRecords.GradingResult(), Learning)
        with open(Outcome.BackupFile, encoding='utf-8') as File:
            Data = json.load(File)
        Rows.append(Row('Normal', 'export a backup after the sample paper and one correction', 'Backup saved; it contains every attempt with its feedback',
                        Pick(Printed, 'Backup saved') + '; ' + str(len(Data['Attempts'])) + ' attempts, the last is attempt ' + str(Data['Attempts'][-1]['AttemptNo'])
                        + ' with ' + str(len(Data['Attempts'][-1]['Feedback']['Points'])) + ' marking-point results'))
        BackupFile = Outcome.BackupFile
        Tests.append({'Name': 'Export a JSON backup', 'Purpose': 'To check that the learning records, including the complete attempt history, can be backed up.', 'Rows': Rows})

        Rows = []
        with open(LearningRecords.DatabasePath, encoding='utf-8') as File:
            Before = File.read()
        Broken = os.path.join(Folder, 'Broken.json')
        with open(Broken, 'w', encoding='utf-8') as File:
            File.write('{"Papers": [')
        WrongType = os.path.join(Folder, 'WrongType.json')
        Data['Papers'][0]['TotalMark'] = '3'
        with open(WrongType, 'w', encoding='utf-8') as File:
            json.dump(Data, File)
        for Kind, Name, Description, Expected in [
                ('Abnormal', 'SampleData/Invalid_Backup.json', 'SampleData/Invalid_Backup.json ("Mistakes" is text, not a list)', 'Import refused with a clear message'),
                ('Abnormal', Broken, 'Broken.json (the JSON stops half way)', 'Import refused: not valid JSON'),
                ('Abnormal', WrongType, 'WrongType.json (TotalMark written as text "3")', 'Import refused, naming the wrong field'),
                ('Abnormal', 'SampleData/Sample_Chemistry_MS.txt', 'a .txt file instead of a .json file', 'Import refused: must be a .json file')]:
            Outcome, Printed = Run(ReportManagement.ManageReports, ['N', 'Y', Name], SQRecords.GradingResult(), SQRecords.LearningData())
            Rows.append(Row(Kind, Description, Expected, '"' + Clean(Outcome.Message) + '" (ImportValid = ' + str(Outcome.ImportValid).upper() + ')'))
        with open(LearningRecords.DatabasePath, encoding='utf-8') as File:
            After = File.read()
        Rows.append(Row('Abnormal', 'compare the records file before and after the four failed imports', 'Saved records not changed', 'Records file unchanged: ' + str(Before == After).upper()))
        Tests.append({'Name': 'Import an invalid JSON file', 'Purpose': 'To check that an invalid import produces a clear message and does not change the saved records.', 'Rows': Rows})

        Rows = []
        os.remove(LearningRecords.DatabasePath)
        Run(LearningRecords.UpdateLearningRecords, [], SQRecords.GradingResult(), SQRecords.RegradeResult())
        Outcome, Printed = Run(ReportManagement.ManageReports, ['N', 'Y', BackupFile], SQRecords.GradingResult(), SQRecords.LearningData())
        Rows.append(Row('Normal', 'the backup made above, imported after the records file was emptied', 'Safety copy of the current records made; backup restored',
                        Pick(Printed, 'first saved as', 'Backup imported')))
        Learning, Printed = Run(LearningRecords.UpdateLearningRecords, ['6'], SQRecords.GradingResult(), SQRecords.RegradeResult())
        Rows.append(Row('Normal', 'the Learning Records Module opens the restored records (integration)', 'The paper score, correction and mistake notebook are back',
                        Pick(Printed, '9701/22/M/J/23', 'Correction attempts saved', 'Mistake notebook entries')))
        Tests.append({'Name': 'Restore a valid backup', 'Purpose': 'To check that a valid JSON backup restores the saved results and correction histories.', 'Rows': Rows})
    return Tests


# ===================================================================== member information

Members = [
    {'Name': 'Fox', 'Module': 'Paper and Document Management Module (with the AI SQ Grader main module and the AI connection)',
     'Files': ['PaperDocumentManagement.py', 'AIConnection.py', 'AISQGrader.py'], 'Tests': FoxTests,
     'Rubric': [
         ('Input', 'User: InputMenuChoice, InputPaperNumber, InputYear, InputAnswerFiles, InputPaperFile (mark scheme and question paper), InputAISettings. File: CheckDocument, ReadSettingsFile (Settings.txt), MakeImageURL (reads a photo). Online data streams: DownloadPaperFile (urllib) and SendChatRequest (the AI service).'),
         ('Collection', 'Parallel arrays SubjectNames / SubjectCodes / SubjectLevels and SeriesNames / SeriesCodes / SeriesLetters; the AnswerFiles list; the 2D SettingsTable; ProviderNames / ProviderURLs; the list of request parts (text and images) in SendChatRequest.'),
         ('Module with parameters', 'CheckDocument(FileName, AllowedTypes) RETURNS STRING; DownloadPaperFile(Paper, PaperType, Download) RETURNS STRING; SendChatRequest(PromptText, ImageURLs, Settings) RETURNS AIResponse; TestConnection(Settings) RETURNS BOOLEAN; ManageDocuments(Download) RETURNS DocumentSet; GetSetting(SettingsTable, Key, DefaultValue).'),
         ('Sequence, selection, iteration', 'InputAnswerFiles: WHILE loop with IF / ELIF / ELSE checks; SendChatRequest: WHILE loop that retries temporary errors, with selection on the HTTP error code; GetFileExtension: WHILE loop with a flag; RunMainMenu.'),
         ('Calls to the module', 'ManageDocuments and TestConnection are called by the main module; CheckDocument by InputAnswerFiles and InputPaperFile; SendChatRequest by Kevin\'s and Caroline\'s modules for every AI request.'),
         ('Output', '"Not added: This file type is not supported ...", "Missing: no mark scheme ...", "Marking cannot begin ...", "The AI service is working." / "The test failed: ...", DisplayDocumentSet, DisplaySettings.'),
         ('Integration', 'The DocumentSet goes to the Document Processing Module (Kevin); AIConnection carries every AI request of Kevin\'s and Caroline\'s modules; the main module calls all six modules in order.')],
     'Video': [
         'Caption "AI SQ Grader - Paper and Document Management (Fox)". Run python AISQGrader.py. If you have an API key, choose 4 (Test the AI connection) and show "The AI service is working."',
         'Choose 1 (Mark a new paper). Input: AS Level, Chemistry, paper 27 (show the error message), then 22, May/June, 2023.',
         'Upload Answers.docx (show "not supported"), then SampleData/Sample_Chemistry_Answers.txt and SampleData/Sample_Answer_Q3.png, then Enter.',
         'Press Enter for the mark scheme: show "Missing: no mark scheme ..." and "The marking session has stopped". Caption "Marking cannot begin without all documents".',
         'Choose 1 again and enter everything (mark scheme SampleData/Sample_Chemistry_MS.txt; Enter for the question paper): "All documents are ready".',
         'Caption "Integration: DocumentSet passed to Kevin\'s Document Processing Module" while its question table appears. Stop recording (under 1 minute).']},
    {'Name': 'Kevin', 'Module': 'Document Processing Module and Grading and Feedback Module',
     'Files': ['DocumentProcessing.py', 'GradingFeedback.py'], 'Tests': KevinTests,
     'Rubric': [
         ('Input', 'Files: mark scheme, question paper and answer files (ReadTextFile, ReadPdfPages, ReadDocumentParts; photos and scanned pages via RenderPdfPage / LoadSourceImages). Online data stream: the AI reads mark scheme pages, the question paper and handwriting, and judges each answer (ConvertMarkSchemeWithAI, TranscribeAnswerImage, RequestAIJudgement).'),
         ('Collection', 'ARRAY OF QuestionItem / MarkingPoint / PointFeedback; the 2D AnswerTable [label, answer, source, confidence]; document parts [source, text, kind]; UncertainMatches; StopWords and UnitWords; key-term lists.'),
         ('Module with parameters', 'ParseMarkScheme(Lines) RETURNS ARRAY OF QuestionItem; TranscribeAnswerImage(ImageURL, SourceRef, Questions, Settings); AddAIQuestions(Questions, Data); AssessMarkingPoints(Question, IsMatched, Settings) RETURNS PointFeedbackSet; CalculateEstimatedMark(Feedback, MarkingRules) RETURNS INTEGER; DiagnoseErrors(Question, Feedback, Settings).'),
         ('Sequence, selection, iteration', 'ParseMarkScheme (FOR loop over lines with an IF / ELIF chain); ConvertMarkSchemeWithAI (nested WHILE loops that send a few pages at a time); GradePaper (FOR loop over questions); IsPointAwarded (recursion).'),
         ('Calls to the module', 'ProcessDocuments and GradePaper are called by the main module; AssessMarkingPoints, CalculateEstimatedMark and TranscribeAnswerImage are also called by the Correction and Regrading Module (Caroline).'),
         ('Output', 'The question table ("found", "found (read from handwriting)", "UNCERTAIN"); the saved SQ mark scheme file; feedback for every marking point; "Not counted: this point depends on P2 ..."; "Likely reason for lost mark ..."; the estimated total.'),
         ('Integration', 'Reads Fox\'s DocumentSet and sends AI requests through Fox\'s AIConnection; returns the GradingResult used by Caroline\'s and Andy\'s modules.')],
     'Video': [
         'Caption "AI SQ Grader - Document Processing and Grading (Kevin)". With an API key for a model that reads images, choose 3 and set it up first (otherwise the offline demo marker is used).',
         'Choose 1 and enter the paper and files quickly (as in Fox\'s video). Caption "Integration: DocumentSet from Fox\'s module".',
         'Show the question table: with the AI, question 3 is "found (read from handwriting)"; offline it is "UNCERTAIN" and flagged for review.',
         'Scroll through the feedback: 2(a) "Not counted: this point depends on P2", 2(b) "Outside the mark scheme", the likely reasons for lost marks.',
         'Show the estimated total. Caption "Output: point-by-point feedback". Answer N to the correction question and stop (under 1 minute).']},
    {'Name': 'Caroline', 'Module': 'Correction and Regrading Module and Learning Records and Progress Module',
     'Files': ['CorrectionRegrading.py', 'LearningRecords.py'], 'Tests': CarolineTests,
     'Rubric': [
         ('Input', 'User: AskYesNo, SelectQuestion, InputCorrectionMethod, InputTypedAnswer, the notebook filter menu. Files: ReadAnswerFile, a photo of the correction (ReadCorrectionPhoto); LoadDatabase reads LearningRecords.json.'),
         ('Collection', 'Attempts (ARRAY OF AttemptRecord); the mistake notebook (ARRAY OF MistakeEntry); parallel arrays TopicNames / TopicWords, TypeNames / TypeCounts and Subjects / PercentTotals / PaperCounts; LearningDatabase lists.'),
         ('Module with parameters', 'RegradeQuestion(Grading, QuestionIndex, NewAnswerText, NewImageRef, Settings) RETURNS RegradeResult; CompareFeedback(OldFeedback, NewFeedback) RETURNS ARRAY OF STRING; FilterMistakes(Database, FilterType, FilterValue); CalculateProgress(Database) RETURNS ProgressSummary; SortCounts(Names, Counts).'),
         ('Sequence, selection, iteration', 'SelectQuestion (WHILE loop); AddRegradeResult (IF / ELIF status rules); SortCounts (bubble sort with NoSwaps and Boundary); FindTopicTags (nested loops); DisplaySubjectAverages.'),
         ('Calls to the module', 'CorrectAndRegrade and UpdateLearningRecords are called by the main module; RegradeQuestion and ReadCorrectionPhoto are called by CorrectAndRegrade.'),
         ('Output', 'The side-by-side comparison, "Well done: the corrected answer gains 2 more mark(s).", the progress dashboard with averages by subject, and the filtered mistake notebook with "To improve" notes.'),
         ('Integration', 'Calls Kevin\'s AssessMarkingPoints, CalculateEstimatedMark and TranscribeAnswerImage to read and reassess a correction; returns LearningData used by Andy\'s module.')],
     'Video': [
         'Caption "AI SQ Grader - Correction, Regrading and Learning Records (Caroline)". Run python AISQGrader.py and mark the sample paper.',
         'Answer Y to the correction question; type 7 (show "Question 7 is not in this paper"), then 2(b); choose 2 and load SampleData/Sample_Corrected_2b.txt (or, with an AI service, choose 3 and upload a photo of a handwritten correction).',
         'Caption "Integration: reassessed by Kevin\'s Grading and Feedback Module". Show "Estimated mark: 0 / 2 -> 2 / 2".',
         'Show the progress dashboard (average by subject); filter the mistake notebook by correction status "Corrected" to show both attempts.',
         'Exit, run the program again, choose 2: the saved records are still there. Stop (under 1 minute).']},
    {'Name': 'Andy', 'Module': 'Report Management Module',
     'Files': ['ReportManagement.py'], 'Tests': AndyTests,
     'Rubric': [
         ('Input', 'User: AskYesNo and the backup file name. File: ImportBackup reads and checks a JSON file; CopyTextFile reads the records file.'),
         ('Collection', 'The report lines (ARRAY OF STRING); pages (ARRAY OF ARRAY OF STRING); the PDF object list and offsets; the 2D field tables PaperFields, MistakeFields, AttemptFields, ErrorFields; the 2D CharacterReplacements table for science symbols.'),
         ('Module with parameters', 'WritePdfReport(ReportLines, FileName) RETURNS BOOLEAN; WrapText(Text, Width) RETURNS ARRAY OF STRING; MakePdfSafe(Text) RETURNS STRING; CheckEntry(Entry, FieldTable, Description) RETURNS STRING; ValidateBackupData(Data) RETURNS STRING; ImportBackup(BackupFileName, DatabaseFileName) RETURNS STRING; ManageReports(Grading, Learning) RETURNS ReportOutcome.'),
         ('Sequence, selection, iteration', 'WritePdfReport (loops that split pages and build objects); WrapText (nested loops and IF); ValidateBackupData (a FOR loop for each section with IF checks); MakePdfSafe (FOR loop over characters with ASC / CHR arithmetic for subscripts).'),
         ('Calls to the module', 'ManageReports is called by the main module; BuildReportLines, WritePdfReport, ExportBackup and ImportBackup are called by ManageReports.'),
         ('Output', 'The PDF report (question text, answer, marking points, diagnoses, attempt history, progress); "PDF report saved: ...", "Backup saved: ...", "Import failed - ... Your saved records have not been changed.", "Backup imported ...".'),
         ('Integration', 'Builds the report from Kevin\'s GradingResult and Caroline\'s LearningData; restores the records file that Caroline\'s module reads.')],
     'Video': [
         'Caption "AI SQ Grader - Report Management (Andy)". Run python AISQGrader.py and mark the sample paper (answer N to the correction question).',
         'Caption "Integration: GradingResult (Kevin) and LearningData (Caroline)". Show "PDF report saved" and open the PDF from the Reports folder.',
         'Answer Y to export a JSON backup: "Backup saved".',
         'Choose 2 in the main menu; import SampleData/Invalid_Backup.json: show "Import failed - ... Your saved records have not been changed."',
         'Import the backup from the Backups folder: "Backup imported". Stop (under 1 minute).']},
]


# ===================================================================== writing the drafts

def MarkdownTable(Tests):
    Lines = ['| Test | Purpose | Test data | Expected outcome | Actual outcome |', '|---|---|---|---|---|']
    for Test in Tests:
        for Index, Item in enumerate(Test['Rows']):
            Name = '**' + Test['Name'] + '**' if Index == 0 else ''
            Purpose = Test['Purpose'] if Index == 0 else ''
            Cells = [Name, Purpose, Item['Data'], Item['Expected'], Item['Actual']]
            Lines.append('| ' + ' | '.join(Cell.replace('|', '\\|').replace('\n', ' ') for Cell in Cells) + ' |')
    return '\n'.join(Lines)


def WriteMarkdown(Member, Tests):
    Text = ['# Part 2 draft - ' + Member['Name'], '',
            '**System/Application:** ' + SystemName + '  ',
            '**Module Name:** ' + Member['Module'] + '  ',
            '**Code to paste:** ' + ', '.join('`' + Name + '`' for Name in Member['Files'])
            + ' (shared records: `SQRecords.py`, written by the whole group)', '',
            'This draft was generated by `Tools/BuildSubmissionDrafts.py`. Every "Actual outcome" below was produced by '
            'running the program; run the tool again after changing the code. Check every row yourself before you submit, '
            'and rewrite the wording in your own words where you wish.', '',
            '## Rubric checklist (Programming techniques and code maintainability)', '',
            '| Requirement | Where it is in the code |', '|---|---|']
    for Requirement, Where in Member['Rubric']:
        Text.append('| ' + Requirement + ' | ' + Where + ' |')
    Text.append('| Comments and identifiers | Header comment with the module\'s job and acknowledgements; `# DECLARE` comments in every subprogram; PascalCase names that describe the data (CAIE style). |')
    Text += ['', '## Test plan', '', MarkdownTable(Tests), '',
             '## Demo video plan (under 1 minute, under 30 MB, no voice; use text captions)', '']
    for Index, Step in enumerate(Member['Video']):
        Text.append(str(Index + 1) + '. ' + Step)
    Text += ['', 'Before recording: run the program from a folder whose path does not contain your name, close other '
             'windows, and do not show your name, ID or school account anywhere in the video.', '']
    with open(os.path.join(OutputFolder, Member['Name'] + '_Part2_Submission.md'), 'w', encoding='utf-8') as File:
        File.write('\n'.join(Text))


def SetCellText(Cell, Text, Bold=False, Size=9):
    from docx.shared import Pt
    Cell.text = ''
    Run = Cell.paragraphs[0].add_run(Text)
    Run.bold = Bold
    Run.font.size = Pt(Size)
    Run.font.name = 'Times New Roman'


def WriteDocx(Member, Tests):
    import docx
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    Document = docx.Document(TemplateFile)
    Paragraphs = Document.paragraphs
    Fill = {'System/Application:': ' ' + SystemName, 'Module Name:': ' ' + Member['Module'],
            'Name & ID:': Member['Name'] + ', ', 'Due date:': DueDate, 'Teacher:': Teacher}
    TestPlanHeading = None
    CodeHeading = None
    TeacherLine = None
    for Paragraph in Paragraphs:
        if Paragraph.text.strip().startswith('Program Code'):
            CodeHeading = Paragraph
        if Paragraph.text.strip().startswith('Teacher:'):
            TeacherLine = Paragraph
        for Label, Value in Fill.items():
            if Paragraph.text.strip().startswith(Label):
                Run = Paragraph.add_run(Value)
                if Paragraph.runs:
                    Run.font.name = Paragraph.runs[0].font.name
        if Paragraph.text.strip().startswith('Test Plan'):
            TestPlanHeading = Paragraph

    # program code: one paragraph per line in the second row of the code table
    Cell = Document.tables[0].rows[1].cells[0]
    for Paragraph in list(Cell.paragraphs)[1:]:
        Paragraph._element.getparent().remove(Paragraph._element)
    First = True
    for FileName in Member['Files']:
        with open(os.path.join(ProgramFolder, FileName), encoding='utf-8') as File:
            CodeLines = ['# ' + '=' * 20 + ' ' + FileName + ' ' + '=' * 20] + File.read().rstrip('\n').split('\n') + ['']
        for Line in CodeLines:
            Paragraph = Cell.paragraphs[0] if First else Cell.add_paragraph()
            First = False
            Paragraph.paragraph_format.space_after = Pt(0)
            Paragraph.paragraph_format.space_before = Pt(0)
            Run = Paragraph.add_run(Line)
            Run.font.name = 'Courier New'
            Run.font.size = Pt(7.5)

    # the cover page ends after "Teacher:"; its empty filler lines are removed and the code
    # starts on a new page, so a long module name cannot push the code onto a later page
    Body = Document.element.body
    Element = TeacherLine._p.getnext()
    while Element is not None and Element is not CodeHeading._p:
        Next = Element.getnext()
        if Element.tag.endswith('}p') and ''.join(Element.itertext()).strip() == '':
            Body.remove(Element)
        Element = Next
    CodeHeading.paragraph_format.page_break_before = True

    # remove the empty lines between the code table and "Test Plan", and after it
    Element = Document.tables[0]._tbl.getnext()
    while Element is not None and Element is not TestPlanHeading._p:
        Next = Element.getnext()
        if Element.tag.endswith('}p') and ''.join(Element.itertext()).strip() == '':
            Body.remove(Element)
        Element = Next
    Element = TestPlanHeading._p.getnext()
    while Element is not None and Element.tag.endswith('}p'):
        Next = Element.getnext()
        if ''.join(Element.itertext()).strip() == '':
            Body.remove(Element)
        Element = Next
    TestPlanHeading.paragraph_format.page_break_before = True

    Intro = Document.add_paragraph()
    IntroRun = Intro.add_run('The test plan for the ' + Member['Module'] + ' includes the following tests. '
                             'The actual outcomes were recorded by running the program.')
    IntroRun.font.name = 'Times New Roman'
    IntroRun.font.size = Pt(11)
    TotalRows = 1 + sum(len(Test['Rows']) for Test in Tests)
    Table = Document.add_table(rows=TotalRows, cols=5)
    Table.style = Document.tables[0].style
    for Column, Heading in enumerate(['Test', 'Purpose', 'Test Data', 'Expected outcome', 'Actual outcome']):
        SetCellText(Table.cell(0, Column), Heading, Bold=True, Size=10)
    RowIndex = 1
    for Test in Tests:
        FirstRow = RowIndex
        for Item in Test['Rows']:
            SetCellText(Table.cell(RowIndex, 2), Item['Data'])
            SetCellText(Table.cell(RowIndex, 3), Item['Expected'])
            SetCellText(Table.cell(RowIndex, 4), Item['Actual'])
            RowIndex = RowIndex + 1
        SetCellText(Table.cell(FirstRow, 0), Test['Name'], Bold=True)
        SetCellText(Table.cell(FirstRow, 1), Test['Purpose'])
        if RowIndex - 1 > FirstRow:
            Table.cell(FirstRow, 0).merge(Table.cell(RowIndex - 1, 0))
            Table.cell(FirstRow, 1).merge(Table.cell(RowIndex - 1, 1))
    TestPlanHeading._p.addnext(Intro._p)
    Intro._p.addnext(Table._tbl)
    Document.save(os.path.join(OutputFolder, Member['Name'] + '_Part2_Submission.docx'))


def Main():
    try:
        import docx  # noqa: F401
        CanWriteDocx = os.path.isfile(TemplateFile)
    except ImportError:
        CanWriteDocx = False
        print('python-docx is not installed (pip install python-docx), so only Markdown drafts are written.')
    for Member in Members:
        print('Running the test plan for', Member['Name'], '...')
        Tests = Member['Tests']()
        WriteMarkdown(Member, Tests)
        if CanWriteDocx:
            WriteDocx(Member, Tests)
    shutil.rmtree(TempRoot, ignore_errors=True)
    print('Drafts written to', OutputFolder)


if __name__ == '__main__':
    Main()
